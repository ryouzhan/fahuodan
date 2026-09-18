import io
import os
import re
import copy
import zipfile
import requests
from datetime import datetime
import pandas as pd
import streamlit as st
import openpyxl
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PILImage

# ===================== 页面全局配置 =====================
st.set_page_config(
    page_title="发货单处理工具 beta",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# ===================== 发票云端默认配置 =====================
INVOICE_CONFIG = {
    "kdocs_webhook": "https://www.kdocs.cn/api/v3/ide/file/cdH0A450EedY/script/V2-6x7HgWruLz4P74YP1PIsVf/sync_task",
    "kdocs_token": "RRZwCZarLOHD4gHKtfSVi",
    "template_urls": {
        "鼎邦物流": "https://gitee.com/zhanliuliang/zll/raw/master/dingbang.xlsx",
        "明日之星": "https://gitee.com/zhanliuliang/zll/raw/master/mingrizhixing.xlsx"
    }
}

# ===================== 简约高级页面样式 (CSS) =====================
st.markdown("""
<style>
    /* 全局背景与字体 */
    .stApp {
        background-color: #F8FAFC !important;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
        color: #1E293B !important;
    }

    /* 隐藏默认水印与顶栏 */
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}

    /* 顶部标题区 */
    .header-box {
        padding: 2rem 0 1rem 0;
        margin-bottom: 0.8rem;
    }
    .header-badge {
        display: inline-block;
        padding: 4px 12px;
        font-size: 0.75rem;
        font-weight: 600;
        letter-spacing: 0.05em;
        text-transform: uppercase;
        background: #EEF2F6;
        color: #475569;
        border-radius: 9999px;
        margin-bottom: 0.5rem;
    }
    .header-title {
        font-size: 2rem;
        font-weight: 700;
        color: #0F172A !important;
        letter-spacing: -0.025em;
        margin: 0;
    }
    .header-subtitle {
        font-size: 0.92rem;
        color: #64748B !important;
        margin-top: 0.35rem;
        font-weight: 400;
    }

    /* 顶部主导航选项卡 Tabs（兼容新旧版本 Streamlit，防深色模式隐形） */
    div[data-testid="stTabs"] [role="tablist"],
    .stTabs [data-baseweb="tab-list"] {
        gap: 12px;
        border-bottom: 1px solid #E2E8F0 !important;
        margin-bottom: 1.5rem;
    }
    div[data-testid="stTabs"] button[role="tab"],
    button[data-testid="stTab"],
    .stTabs [data-baseweb="tab"] {
        padding: 10px 20px !important;
        font-size: 1rem !important;
        background: transparent !important;
        border-radius: 8px 8px 0 0 !important;
        border-bottom: 2.5px solid transparent !important;
    }
    /* 未选中 Tab：强制深灰文本 */
    div[data-testid="stTabs"] button[role="tab"] *,
    button[data-testid="stTab"] *,
    .stTabs [data-baseweb="tab"] * {
        color: #64748B !important;
        font-weight: 500 !important;
    }
    /* 选中 Tab：强制深黑文本与底部高亮横线 */
    div[data-testid="stTabs"] button[role="tab"][aria-selected="true"],
    button[data-testid="stTab"][aria-selected="true"],
    .stTabs [data-baseweb="tab"][aria-selected="true"] {
        border-bottom: 2.5px solid #0F172A !important;
    }
    div[data-testid="stTabs"] button[role="tab"][aria-selected="true"] *,
    button[data-testid="stTab"][aria-selected="true"] *,
    .stTabs [data-baseweb="tab"][aria-selected="true"] * {
        color: #0F172A !important;
        font-weight: 700 !important;
    }

    /* 上传框白底样式 */
    [data-testid="stFileUploader"] {
        background: transparent !important;
    }
    [data-testid="stFileUploader"] section {
        background-color: #FFFFFF !important;
        border: 1.5px dashed #CBD5E1 !important;
        border-radius: 12px !important;
        padding: 1.4rem 1rem !important;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.02) !important;
        transition: all 0.2s ease !important;
    }
    [data-testid="stFileUploader"] section:hover {
        border-color: #3B82F6 !important;
        background-color: #F8FAFC !important;
    }
    [data-testid="stFileUploader"] section * {
        color: #475569 !important;
    }
    [data-testid="stFileUploader"] section button {
        background-color: #F1F5F9 !important;
        color: #0F172A !important;
        border: 1px solid #E2E8F0 !important;
        border-radius: 8px !important;
    }
    [data-testid="stFileUploader"] label,
    [data-testid="stFileUploader"] label p {
        color: #1E293B !important;
        font-size: 0.95rem !important;
        font-weight: 600 !important;
    }

    /* 下拉选择框容器与文字强制深色可见 */
    [data-testid="stSelectbox"] label p {
        color: #1E293B !important;
        font-size: 0.95rem !important;
        font-weight: 600 !important;
    }
    [data-testid="stSelectbox"] div[data-baseweb="select"] {
        background-color: #FFFFFF !important;
        border-radius: 8px !important;
    }
    [data-testid="stSelectbox"] div[data-baseweb="select"] * {
        color: #1E293B !important;
    }

    /* KPI 卡片网格 */
    .metric-container {
        display: grid;
        grid-template-columns: repeat(4, 1fr);
        gap: 1.2rem;
        margin: 1.5rem 0;
    }
    .metric-card {
        background: #FFFFFF;
        padding: 1.3rem 1.2rem;
        border-radius: 12px;
        border: 1px solid #E2E8F0;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.02);
        transition: transform 0.15s ease;
    }
    .metric-card:hover {
        transform: translateY(-2px);
        box-shadow: 0 4px 12px rgba(0, 0, 0, 0.05);
    }
    .metric-title {
        font-size: 0.8rem;
        font-weight: 500;
        color: #64748B;
        margin-bottom: 0.35rem;
        text-transform: uppercase;
    }
    .metric-num {
        font-size: 1.7rem;
        font-weight: 700;
        color: #0F172A;
        line-height: 1.2;
    }
    .metric-unit {
        font-size: 0.85rem;
        font-weight: 500;
        color: #94A3B8;
        margin-left: 0.2rem;
    }

    /* 按钮基础风格 */
    div.stButton > button:first-child {
        background-color: #0F172A !important;
        color: #FFFFFF !important;
        border-radius: 8px !important;
        border: none !important;
        padding: 0.7rem 1.8rem !important;
        font-size: 1rem !important;
        font-weight: 600 !important;
        box-shadow: 0 3px 8px rgba(15, 23, 42, 0.15) !important;
    }
    div.stButton > button:first-child:hover {
        background-color: #1E293B !important;
        transform: translateY(-1px);
    }
    .stDownloadButton button {
        background-color: #0F172A !important;
        color: #FFFFFF !important;
        border-radius: 8px !important;
        border: none !important;
        padding: 0.75rem 2rem !important;
        font-size: 1rem !important;
        font-weight: 600 !important;
        box-shadow: 0 3px 10px rgba(15, 23, 42, 0.15) !important;
    }
    .stDownloadButton button:hover {
        background-color: #1E293B !important;
        transform: translateY(-1px);
    }
</style>
""", unsafe_allow_html=True)

# =====================================================================
# 模块 1：发货单处理业务逻辑
# =====================================================================
def beautify_excel(file_stream, col_width=15, row_height=20):
    """美化Excel表格，并根据<货件编号>设置整行单元格颜色，处理所有相关Sheet"""
    try:
        wb = load_workbook(file_stream)
        
        font = Font(name='Calibri', size=11, bold=True)
        alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        thin = Side(border_style="thin", color="000000")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        colors = [
            "FFCCCC", "CCFFCC", "CCCCFF", "FFFFCC", "CCFFFF", "FFCCFF",
            "FF9999", "99FF99", "9999FF", "FFCC99", "99FFCC", "CCCC99"
        ]

        for sheet_name in ['详细数据', '汇总结果']:
            if sheet_name not in wb.sheetnames:
                continue

            ws = wb[sheet_name]
            for col in ws.columns:
                column_letter = col[0].column_letter
                ws.column_dimensions[column_letter].width = col_width

            for row in ws.iter_rows():
                if not row[0].row:
                    continue
                ws.row_dimensions[row[0].row].height = row_height
                for cell in row:
                    if cell.value is None:
                        continue
                    cell.alignment = alignment
                    cell.border = border
                    if cell.row == 1:
                        cell.font = font

            if sheet_name == '详细数据':
                header_row = [cell.value for cell in ws[1]]
                if "货件编号" not in header_row:
                    continue
                item_code_col_idx = header_row.index("货件编号") + 1

                item_codes = {}
                for row in ws.iter_rows(min_row=2, min_col=item_code_col_idx, max_col=item_code_col_idx):
                    val = row[0].value
                    if val and val not in item_codes:
                        item_codes[val] = None

                unique_codes = list(item_codes.keys())
                for idx, code in enumerate(unique_codes):
                    item_codes[code] = colors[idx % len(colors)]

                for row in ws.iter_rows(min_row=2):
                    code_cell = row[item_code_col_idx - 1]
                    code_val = code_cell.value
                    if code_val in item_codes:
                        fill = PatternFill(start_color=item_codes[code_val],
                                           end_color=item_codes[code_val],
                                           fill_type="solid")
                        for cell in row:
                            cell.fill = fill

        output_stream = io.BytesIO()
        wb.save(output_stream)
        output_stream.seek(0)
        return output_stream

    except Exception as e:
        raise Exception(f"美化出错: {str(e)}")

def load_and_merge_data(uploaded_file):
    """加载 发货单详情 + 装箱信息 并合并"""
    try:
        xl_file = pd.ExcelFile(uploaded_file)
        sheet_names = xl_file.sheet_names

        required_sheets = ['发货单详情', '装箱信息']
        for sheet in required_sheets:
            if sheet not in sheet_names:
                raise KeyError(f"文件缺少必须工作表：{sheet}")

        df_detail = pd.read_excel(uploaded_file, sheet_name='发货单详情')
        df_pack = pd.read_excel(uploaded_file, sheet_name='装箱信息')

        df_detail.rename(columns={
            '货件编号(发货商品)': '货件编号',
            '发货单号(基本信息)': '发货单号',
            'SKU(发货商品)': 'SKU',
            '发货量(发货商品)': '发货量',
            '品名(发货商品)': '品名',
            '商品图片(发货商品)': '商品图片',
            '商品品牌': '供应商'
        }, inplace=True)

        df_pack.rename(columns={
            '货件编号(装箱信息)': '货件编号',
            '发货单号(装箱信息)': '发货单号',
            'SKU(装箱信息)': 'SKU',
            '发货量(装箱信息)': '发货量',
            'ReferenceId(装箱信息)': 'ReferenceId'
        }, inplace=True)

        merged = pd.merge(
            df_detail, df_pack,
            on=['货件编号', '发货单号', 'SKU', '发货量'],
            how='left'
        )

        keep_cols = [
            'SKU', '品名', '商品图片', '发货量', '箱数', '箱规单箱数量', '物流中心编码',
            '供应商', '货件编号', 'ReferenceId', '箱号', '总箱数',
            '外箱重量(kg)', '外箱总重量(kg)', '外箱长(cm)', '外箱宽(cm)', '外箱高(cm)',
            '外箱体积(m³)', '外箱总体积(m³)', '外箱总体积重(kg)', '创建时间', '发货时间', '物流商'
        ]
        exist_cols = [c for c in keep_cols if c in merged.columns]
        merged = merged[exist_cols]

        merged.rename(columns={
            '箱规单箱数量': '单箱数量',
            '总箱数': '总箱数编号'
        }, inplace=True)

        return merged
    except Exception as e:
        raise Exception(f"数据加载合并失败：{str(e)}")

def perform_summary(df):
    """按 物流中心编码 + 货件编号 + ReferenceId 汇总"""
    df_copy = df.copy()
    df_copy['物流中心编码'] = df_copy['物流中心编码'].fillna('未知地址')
    df_copy['货件编号'] = df_copy['货件编号'].fillna('未知编号')
    df_copy['ReferenceId'] = df_copy['ReferenceId'].fillna('无ReferenceId')

    summary = df_copy.groupby(['物流中心编码', '货件编号', 'ReferenceId'], dropna=False).agg(
        品名=('品名', 'first'),
        供应商=('供应商', 'first'),
        总箱数=('箱数', 'sum'),
        外箱总重量=('外箱总重量(kg)', 'sum'),
        外箱总体积=('外箱总体积(m³)', 'sum')
    ).reset_index()

    summary['外箱总体积'] = summary['外箱总体积'].fillna(0)
    summary['外箱总体积重(kg)'] = (summary['外箱总体积'] * 167).round(2)

    return summary

def save_excel(df, summary_df):
    """保存为两个Sheet：详细数据 + 汇总结果"""
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='详细数据', index=False)
        summary_df.to_excel(writer, sheet_name='汇总结果', index=False)
    output.seek(0)
    return output

# =====================================================================
# 模块 2：物流发票处理业务逻辑
# =====================================================================
def fetch_kdocs_po_data(webhook_url, api_token):
    url = webhook_url.strip()
    token = api_token.strip()
    if not url or not token:
        raise ValueError("未配置有效的金山文档 Webhook 或 Token！")

    headers = {"AirScript-Token": token, "Content-Type": "application/json"}
    resp = requests.post(url, json={"Context": {"argv": {}}}, headers=headers, timeout=25)
    if resp.status_code != 200:
        raise ValueError(f"请求金山文档失败 (HTTP {resp.status_code}): {resp.text}")

    res_json = resp.json()
    data = None
    if isinstance(res_json, dict):
        if "data" in res_json and isinstance(res_json["data"], dict) and "result" in res_json["data"]:
            data = res_json["data"]["result"]
        elif "result" in res_json:
            data = res_json["result"]
        elif "data" in res_json and isinstance(res_json["data"], list):
            data = res_json["data"]
    elif isinstance(res_json, list):
        data = res_json

    if not data or not isinstance(data, list):
        raise ValueError(f"金山文档未返回有效数据列表，返回内容: {res_json}")

    return pd.DataFrame(data)

def fetch_template_bytes(template_url, carrier_name):
    url = str(template_url).strip() if template_url else ""
    if "/blob/" in url:
        url = re.sub(r'/blob/([^/]+)/', r'/raw/\1/', url)

    if url and url.startswith(("http://", "https://")):
        try:
            resp = requests.get(url, timeout=20)
            if resp.status_code == 200 and len(resp.content) > 1000 and not resp.content.startswith(b"<!DOCTYPE"):
                return resp.content
        except Exception:
            pass

    app_dir = os.path.dirname(os.path.abspath(__file__)) if '__file__' in globals() else os.getcwd()
    local_candidates = [
        os.path.join(app_dir, f"{carrier_name}.xlsx"),
        os.path.join(app_dir, f"{carrier_name}发票.xlsx"),
    ]
    if "鼎邦" in carrier_name:
        local_candidates.append(os.path.join(app_dir, "dingbang.xlsx"))
    elif "明日" in carrier_name:
        local_candidates.append(os.path.join(app_dir, "mingrizhixing.xlsx"))

    for loc_file in local_candidates:
        if loc_file and os.path.isfile(loc_file):
            with open(loc_file, 'rb') as f:
                return f.read()

    return None

class InvoiceProcessor:
    @staticmethod
    def clean_sku(val):
        if val is None or pd.isna(val):
            return ""
        s = str(val).strip()
        if s.endswith('.0') and s[:-2].isdigit():
            s = s[:-2]
        return s.strip("'\"\t\r\n \u3000")

    @staticmethod
    def normalize_carrier_name(raw_name):
        if not raw_name or pd.isna(raw_name):
            return ""
        name = str(raw_name).strip()
        name = re.sub(r'[\r\n\t]+', '', name)
        if '鼎邦' in name:
            return '鼎邦物流'
        if '明日' in name:
            return '明日之星'
        if '易达' in name:
            return '易达物流'
        return name

    @staticmethod
    def load_summary_sheet_data(file_obj):
        ref_mapping = {}
        carrier_mapping = {}
        summary_name = None
        try:
            excel_obj = pd.ExcelFile(file_obj)
            sheet_names = excel_obj.sheet_names
            target_sheet = None
            for s in sheet_names:
                if '汇总结果' in str(s) or '汇总' in str(s):
                    target_sheet = s
                    break
            if target_sheet is None and len(sheet_names) >= 2:
                target_sheet = sheet_names

            if target_sheet is not None:
                summary_name = str(target_sheet)
                df_summary = pd.read_excel(excel_obj, sheet_name=target_sheet)

                fba_col, ref_col, carrier_col = None, None, None
                for c in df_summary.columns:
                    c_clean = str(c).strip().lower().replace(' ', '').replace('_', '').replace('*', '')
                    if not fba_col and any(cand in c_clean for cand in ['货件编号', 'fba编号', 'shipmentid', 'fbano', '货件号']):
                        fba_col = c
                    elif not ref_col and any(cand in c_clean for cand in ['referenceid', 'reference', '货物追踪编号', '参考号', '追踪编号', '追踪号']):
                        ref_col = c

                for cand in ['物流商', '承运商', '物流公司', '渠道', '发货物流', '物流']:
                    for c in df_summary.columns:
                        c_clean = str(c).strip().replace('*', '')
                        if any(ex in c_clean for ex in ['中心', '编码', '仓', '代码', 'code', '单号', '体积', '重', '费']):
                            continue
                        if cand in c_clean:
                            carrier_col = c
                            break
                    if carrier_col:
                        break

                if fba_col:
                    for _, row in df_summary.iterrows():
                        fba_val = str(row[fba_col]).strip() if pd.notna(row[fba_col]) else ""
                        if not fba_val or fba_val.lower() == 'nan':
                            continue
                        if ref_col and pd.notna(row[ref_col]):
                            ref_val = str(row[ref_col]).strip()
                            if ref_val and ref_val.lower() != 'nan':
                                ref_mapping[fba_val] = ref_val
                                ref_mapping[fba_val.upper()] = ref_val
                        if carrier_col and pd.notna(row[carrier_col]):
                            car_val = str(row[carrier_col]).strip()
                            if car_val and car_val.lower() != 'nan':
                                car_clean = InvoiceProcessor.normalize_carrier_name(car_val)
                                carrier_mapping[fba_val] = car_clean
                                carrier_mapping[fba_val.upper()] = car_clean
        except Exception:
            pass
        return ref_mapping, carrier_mapping, summary_name

    @staticmethod
    def clean_name(name):
        return re.sub(r'[\\/*?:"<>|]', '_', str(name).strip())

    @staticmethod
    def clean_warehouse_code(code):
        return re.sub(r'[\(（][^\)）]*[\)）]', '', str(code)).strip() if code else ""

    @staticmethod
    def get_dict_val(data_dict, candidate_keys):
        for cand in candidate_keys:
            cand_clean = str(cand).strip().lower()
            for k, val in data_dict.items():
                if cand_clean == str(k).strip().lower():
                    if pd.notna(val) and str(val).strip() != "":
                        return val
        return ""

    @staticmethod
    def parse_box_ranges(box_str):
        if not box_str:
            return None
        s = str(box_str).strip()
        s = re.sub(r'[,;，；、\s]+', ',', s)
        boxes = set()
        has_number = False
        for part in s.split(','):
            part = part.strip()
            if not part:
                continue
            m_range = re.search(r'(\d+)\s*[-~至到_–—－~～]\s*(\d+)', part)
            if m_range:
                start, end = int(m_range.group(1)), int(m_range.group(2))
                boxes.update(range(min(start, end), max(start, end) + 1))
                has_number = True
            else:
                m_single = re.findall(r'\b\d+\b', part)
                if m_single:
                    for num in m_single:
                        boxes.add(int(num))
                    has_number = True
        return boxes if has_number else None

    @staticmethod
    def extract_row_boxes(s_dict, cur_box):
        start_box_val = InvoiceProcessor.get_dict_val(s_dict, ['起始箱号', '开始箱号', '起始箱'])
        end_box_val = InvoiceProcessor.get_dict_val(s_dict, ['截止箱号', '结束箱号', '终止箱号', '截止箱'])
        if start_box_val and end_box_val:
            try:
                s_int = int(re.search(r'\d+', str(start_box_val)).group(0))
                e_int = int(re.search(r'\d+', str(end_box_val)).group(0))
                return set(range(min(s_int, e_int), max(s_int, e_int) + 1)), max(s_int, e_int) - min(s_int, e_int) + 1
            except Exception:
                pass

        box_no_val = InvoiceProcessor.get_dict_val(s_dict, ['箱号', '外箱箱号', '箱号编号', 'CTN NO', 'Carton No'])
        box_cnt_val = InvoiceProcessor.get_dict_val(s_dict, ['箱数', '外箱数', '总箱数'])

        if box_no_val and re.search(r'(\d+)\s*[-~至到_–—－~～]\s*(\d+)', str(box_no_val)):
            b = InvoiceProcessor.parse_box_ranges(box_no_val)
            if b:
                return b, len(b)

        count_int = 1
        if box_cnt_val:
            try:
                count_int = int(float(str(box_cnt_val).strip()))
            except Exception:
                pass

        if box_no_val and str(box_no_val).strip().isdigit() and (not box_cnt_val or str(box_no_val).strip() != str(box_cnt_val).strip()):
            start_num = int(str(box_no_val).strip())
            return set(range(start_num, start_num + count_int)), count_int

        return set(range(cur_box, cur_box + count_int)), count_int

    @staticmethod
    def format_box_ranges(boxes_set):
        if not boxes_set:
            return ""
        sorted_boxes = sorted(list(boxes_set))
        if len(sorted_boxes) == 1:
            return f"{sorted_boxes[0]}"
        if sorted_boxes[-1] - sorted_boxes[0] + 1 == len(sorted_boxes):
            return f"{sorted_boxes[0]}-{sorted_boxes[-1]}"
        ranges = []
        start = sorted_boxes[0]
        end = sorted_boxes[0]
        for n in sorted_boxes[1:]:
            if n == end + 1:
                end = n
            else:
                ranges.append(f"{start}" if start == end else f"{start}-{end}")
                start = n
                end = n
        ranges.append(f"{start}" if start == end else f"{start}-{end}")
        return ",".join(ranges)

    @staticmethod
    def format_dimensions(length, width, height):
        def clean_val(v):
            if v is None:
                return ""
            s = str(v).strip()
            if not s or s.lower() == "nan":
                return ""
            try:
                fv = float(s)
                return f"{int(fv)}" if fv.is_integer() else f"{fv}"
            except Exception:
                return s
        l_str = clean_val(length)
        w_str = clean_val(width)
        h_str = clean_val(height)
        if l_str and w_str and h_str:
            return f"{l_str}*{w_str}*{h_str}"
        return l_str or w_str or h_str or ""

# =====================================================================
# 页面顶部 Header
# =====================================================================
st.markdown("""
<div class="header-box">
    <div class="header-badge">Logistics All-in-One</div>
    <h1 class="header-title">发货单处理工具</h1>
    <p class="header-subtitle">by Ryou</p>
</div>
""", unsafe_allow_html=True)

# 顶部主导航标签页
tab_shipment, tab_invoice = st.tabs(["📦 发货单合并汇总工具", "📑 物流发票批量生成工具"])


# =====================================================================
# 工具 1：发货单合并汇总
# =====================================================================
with tab_shipment:
    uploaded_ship_file = st.file_uploader(
        "上传发货单原始 Excel 文件",
        type=["xlsx", "xls"],
        help="需包含「发货单详情」与「装箱信息」两张工作表",
        key="ship_uploader"
    )

    if uploaded_ship_file is not None:
        try:
            with st.spinner("正在处理发货单数据并套用美化格式..."):
                merged_df = load_and_merge_data(uploaded_ship_file)
                summary_df = perform_summary(merged_df)
                raw_excel = save_excel(merged_df, summary_df)
                excel_bytes = beautify_excel(raw_excel)

            # 4 个核心 KPI 指标卡
            total_boxes = int(summary_df['总箱数'].sum())
            total_weight = f"{summary_df['外箱总重量'].sum():,.2f}"
            total_vol = f"{summary_df['外箱总体积'].sum():,.2f}"
            total_vol_weight = f"{summary_df['外箱总体积重(kg)'].sum():,.2f}"

            st.markdown(f"""
            <div class="metric-container">
                <div class="metric-card">
                    <div class="metric-title">总装箱量</div>
                    <div class="metric-num">{total_boxes:,}<span class="metric-unit">箱</span></div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">实重合计</div>
                    <div class="metric-num">{total_weight}<span class="metric-unit">kg</span></div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">外箱总体积</div>
                    <div class="metric-num">{total_vol}<span class="metric-unit">m³</span></div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">总体积重 (1:167)</div>
                    <div class="metric-num">{total_vol_weight}<span class="metric-unit">kg</span></div>
                </div>
            </div>
            """, unsafe_allow_html=True)

            st.write("")

            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            download_filename = f"发货单处理结果_{timestamp}.xlsx"

            # 居中导出按钮
            col_left, col_btn, col_right = st.columns(3)
            with col_btn:
                st.download_button(
                    label="⬇️ 导出处理后的 Excel 文件",
                    data=excel_bytes,
                    file_name=download_filename,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

        except Exception as e:
            st.error(f"处理发生异常：{str(e)}")


# =====================================================================
# 工具 2：物流发票批量生成
# =====================================================================
with tab_invoice:
    carrier_option = st.selectbox(
        "📦 选择物流模板模式",
        options=["鼎邦物流", "明日之星", "【自动识别】按表格物流列"],
        index=0,
        help="选择生成发票所使用的报关模板",
        key="inv_carrier_sel"
    )

    uploaded_inv_file = st.file_uploader(
        "上传发货单处理结果文件",
        type=["xlsx", "xls"],
        key="inv_uploader"
    )

    if uploaded_inv_file is not None:
        if st.button("🚀 开始批量生成发票", use_container_width=True, key="btn_run_inv"):
            progress_bar = st.progress(0)
            status_text = st.empty()

            try:
                status_text.text("【1/4】正在连接金山文档云端，同步采购单主数据...")
                df_po = fetch_kdocs_po_data(INVOICE_CONFIG["kdocs_webhook"], INVOICE_CONFIG["kdocs_token"])
                progress_bar.progress(25)

                po_dict = {}
                sku_col = next((c for c in df_po.columns if 'sku' in str(c).strip().lower()), None)
                if not sku_col:
                    raise ValueError("采购单中未找到 SKU 列")
                for _, row in df_po.iterrows():
                    sku_val = InvoiceProcessor.clean_sku(row.get(sku_col))
                    if sku_val:
                        po_dict[sku_val] = row.to_dict()
                        po_dict[sku_val.upper()] = row.to_dict()
                        po_dict[sku_val.lower()] = row.to_dict()

                status_text.text("【2/4】正在解析发货单明细与汇总数据...")
                uploaded_inv_file.seek(0)
                df_ship = pd.read_excel(uploaded_inv_file)
                
                ship_sku_col = next((c for c in df_ship.columns if 'sku' in str(c).strip().lower()), 'SKU')
                ship_id_col = next((c for c in df_ship.columns if any(k in str(c) for k in ['货件编号', 'FBA货件编号', 'ShipmentID'])), '货件编号')

                if ship_id_col not in df_ship.columns:
                    raise ValueError(f"发货单中未找到【货件编号】列，现有列: {list(df_ship.columns)}")

                uploaded_inv_file.seek(0)
                summary_ref_mapping, summary_carrier_mapping, _ = InvoiceProcessor.load_summary_sheet_data(uploaded_inv_file)

                valid_rows = [row for _, row in df_ship.iterrows() if pd.notna(row.get(ship_id_col)) and str(row.get(ship_id_col)).strip() not in ('', 'nan')]
                df_valid_ship = pd.DataFrame(valid_rows)
                groups = df_valid_ship.groupby(ship_id_col)
                total_groups = len(groups)

                progress_bar.progress(40)
                status_text.text(f"【3/4】共识别到 {total_groups} 个独立货件，正在匹配生成发票...")

                template_cache = {}
                generated_zip_files = {}
                missing_records, skipped_records = [], []
                center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                thin_border = Border(left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'), top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3'))

                current_idx = 0
                for fba_shipment_id, group_df in groups:
                    current_idx += 1
                    first_row_data = group_df.iloc[0].to_dict()

                    # 物流识别
                    current_carrier = ""
                    if carrier_option != "【自动识别】按表格物流列":
                        current_carrier = InvoiceProcessor.normalize_carrier_name(carrier_option)
                    else:
                        current_carrier = summary_carrier_mapping.get(fba_shipment_id) or ""
                        if not current_carrier:
                            for col_k in ['物流', '物流商', '承运商', '物流公司', '渠道', '发货物流']:
                                v = first_row_data.get(col_k)
                                if pd.notna(v) and str(v).strip():
                                    current_carrier = InvoiceProcessor.normalize_carrier_name(v)
                                    break
                    if not current_carrier:
                        current_carrier = "未知物流"

                    if current_carrier not in template_cache:
                        tpl_url = INVOICE_CONFIG["template_urls"].get(current_carrier, "")
                        template_cache[current_carrier] = fetch_template_bytes(tpl_url, current_carrier)

                    template_bytes = template_cache.get(current_carrier)
                    if not template_bytes:
                        skipped_records.append({'货件编号': fba_shipment_id, '识别物流': current_carrier, '原因': '未配置发票模板'})
                        continue

                    # 提取账号与仓码
                    account_name = "未知"
                    for _, s_row in group_df.iterrows():
                        s_sku = InvoiceProcessor.clean_sku(s_row.get(ship_sku_col, ''))
                        if len(s_sku) >= 2:
                            account_name = s_sku[:2].upper()
                            break

                    clean_acc = InvoiceProcessor.clean_name(account_name)
                    clean_fba = InvoiceProcessor.clean_name(fba_shipment_id)
                    clean_fba_wh_code = InvoiceProcessor.clean_warehouse_code(InvoiceProcessor.get_dict_val(first_row_data, ['物流中心编码', '仓码', '仓库编码']))

                    val_ref_id = summary_ref_mapping.get(fba_shipment_id) or ""
                    if not val_ref_id:
                        for _, s_row in group_df.iterrows():
                            r_val = InvoiceProcessor.get_dict_val(s_row.to_dict(), ['ReferenceId', 'Reference ID', 'Reference', '参考号'])
                            if r_val:
                                val_ref_id = str(r_val).strip()
                                break

                    raw_items = []
                    all_shipment_boxes = set()
                    cur_box = 1

                    for _, s_row in group_df.iterrows():
                        s_dict = s_row.to_dict()
                        s_sku = InvoiceProcessor.clean_sku(s_dict.get(ship_sku_col, ''))
                        po_item = po_dict.get(s_sku) or {}

                        if not po_item:
                            missing_records.append({'货件编号': fba_shipment_id, '未匹配SKU': s_sku, '品名': s_dict.get('品名', '')})

                        parsed_boxes, row_box_cnt = InvoiceProcessor.extract_row_boxes(s_dict, cur_box)
                        if parsed_boxes:
                            all_shipment_boxes.update(parsed_boxes)
                            cur_box = max(cur_box, max(parsed_boxes) + 1)

                        val_len = InvoiceProcessor.get_dict_val(s_dict, ['外箱长(cm)', '外箱长', '长(cm)', '长'])
                        val_width = InvoiceProcessor.get_dict_val(s_dict, ['外箱宽(cm)', '外箱宽', '宽(cm)', '宽'])
                        val_height = InvoiceProcessor.get_dict_val(s_dict, ['外箱高(cm)', '外箱高', '高(cm)', '高'])
                        val_vol_dim = InvoiceProcessor.format_dimensions(val_len, val_width, val_height)

                        raw_items.append({
                            'sku': s_sku,
                            'box_count': row_box_cnt,
                            'val_box_num': InvoiceProcessor.format_box_ranges(parsed_boxes),
                            'parsed_boxes': parsed_boxes,
                            'val_fba_no': fba_shipment_id,
                            'val_ref_id': val_ref_id,
                            'val_name_cn': InvoiceProcessor.get_dict_val(po_item, ['品名', '中文品名']) or InvoiceProcessor.get_dict_val(s_dict, ['品名']),
                            'val_name_en': InvoiceProcessor.get_dict_val(po_item, ['英文品名', '英文名']),
                            'val_qty_per_box': InvoiceProcessor.get_dict_val(s_dict, ['单箱数量', 'PCS/CTN']),
                            'val_unit_price': InvoiceProcessor.get_dict_val(po_item, ['申报单价', '单价']),
                            'val_material': InvoiceProcessor.get_dict_val(po_item, ['材质中/英*', '材质中/英', '材质']),
                            'val_usage': InvoiceProcessor.get_dict_val(po_item, ['用途中/英*', '用途中/英', '用途']),
                            'val_hs_code': InvoiceProcessor.get_dict_val(po_item, ['HS编码', '海关编码']),
                            'val_weight': InvoiceProcessor.get_dict_val(s_dict, ['外箱重量(kg)', '外箱重量', '单箱重']),
                            'val_len': val_len, 'val_width': val_width, 'val_height': val_height,
                            'val_vol_dim': val_vol_dim,
                            'val_img_url': InvoiceProcessor.get_dict_val(s_dict, ['商品图片', '图片'])
                        })

                    # 箱号合并
                    raw_items.sort(key=lambda it: min(it['parsed_boxes']) if it.get('parsed_boxes') else 999999)
                    merged_items = []
                    for item in raw_items:
                        cn, en = str(item['val_name_cn']).strip().lower(), str(item['val_name_en']).strip().lower()
                        curr_start = min(item['parsed_boxes']) if item.get('parsed_boxes') else None

                        can_merge = False
                        if merged_items:
                            prev = merged_items[-1]
                            prev_cn, prev_en = str(prev['val_name_cn']).strip().lower(), str(prev['val_name_en']).strip().lower()
                            if (cn, en) == (prev_cn, prev_en):
                                prev_end = max(prev['parsed_boxes']) if prev.get('parsed_boxes') else None
                                if prev_end is not None and curr_start is not None and curr_start == prev_end + 1:
                                    can_merge = True

                        if can_merge:
                            prev = merged_items[-1]
                            prev['box_count'] += item['box_count']
                            if prev['parsed_boxes'] and item['parsed_boxes']:
                                prev['parsed_boxes'].update(item['parsed_boxes'])
                                prev['val_box_num'] = InvoiceProcessor.format_box_ranges(prev['parsed_boxes'])
                        else:
                            merged_items.append(copy.copy(item))

                    final_items = merged_items
                    total_boxes = sum(it['box_count'] for it in raw_items) or len(all_shipment_boxes) or len(group_df)

                    # 写入 Workbook
                    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
                    ws = wb.active

                    header_filled = {}
                    if current_carrier == "明日之星":
                        for r in range(1, min(15, ws.max_row + 1)):
                            for c in range(1, min(20, ws.max_column + 1)):
                                cell_val = str(ws.cell(row=r, column=c).value or "").strip()
                                if any(k in cell_val for k in ['批次跟踪号', 'Ref.']) and '批次' not in header_filled:
                                    ws.cell(row=r, column=c + 1, value=fba_shipment_id).alignment = center_alignment
                                    header_filled['批次'] = True
                                elif '箱数/件数' in cell_val and '总箱数' not in header_filled:
                                    ws.cell(row=r, column=c + 1, value=total_boxes).alignment = center_alignment
                                    header_filled['总箱数'] = True
                                elif any(k in cell_val for k in ['Consignee', '收件人:']) and 'Consignee' not in header_filled:
                                    ws.cell(row=r, column=c + 1, value=clean_fba_wh_code).alignment = center_alignment
                                    header_filled['Consignee'] = True
                    else:
                        for r in range(1, min(15, ws.max_row + 1)):
                            for c in range(1, min(20, ws.max_column + 1)):
                                cell_val = str(ws.cell(row=r, column=c).value or "").strip()
                                if '客户单号' in cell_val and '客户单号' not in header_filled:
                                    ws.cell(row=r, column=c + 1, value=fba_shipment_id).alignment = center_alignment
                                    header_filled['客户单号'] = True
                                elif ('FBA仓码' in cell_val or '仓码' in cell_val) and 'FBA仓码' not in header_filled:
                                    ws.cell(row=r, column=c + 1, value=clean_fba_wh_code).alignment = center_alignment
                                    header_filled['FBA仓码'] = True

                    header_row_idx = None
                    col_mapping = {}
                    for r in range(1, min(25, ws.max_row + 1)):
                        row_vals = [str(ws.cell(row=r, column=c).value or "").strip() for c in range(1, ws.max_column + 1)]
                        if sum(1 for v in row_vals if any(k in v for k in ['箱数', '中文品名', '英文品名', '单箱数量', '单价', '海关编码'])) >= 3:
                            header_row_idx = r
                            for c in range(1, ws.max_column + 1):
                                cv = str(ws.cell(row=r, column=c).value or "").strip()
                                if cv:
                                    col_mapping[cv] = c
                            break

                    if not header_row_idx:
                        continue

                    def get_col(candidates):
                        for cand in candidates:
                            for h_name, col_i in col_mapping.items():
                                if cand in h_name:
                                    return col_i
                        return None

                    col_fba = get_col(['FBA编号', '货件编号'])
                    col_ref = get_col(['Reference ID', 'ReferenceId', '参考号'])
                    col_vol = get_col(['材积CM', '长*宽*高', '材积'])
                    col_box = get_col(['箱数', '箱号'])
                    col_wt = get_col(['单件毛重', '单箱重', '外箱重量'])
                    col_hs = get_col(['海关编码', 'HSCODE', 'HS编码'])
                    col_cn = get_col(['中文品名', '中文名', '品名'])
                    col_en = get_col(['英文品名', '英文名'])
                    col_qty = get_col(['单箱数量', 'PCS/CTN', '每箱数量'])
                    col_price = get_col(['单价', '单价USD', '申报单价'])
                    col_mat = get_col(['材质中/英*', '材质中/英', '材质'])
                    col_use = get_col(['用途中/英*', '用途中/英', '用途'])
                    col_img = get_col(['产品图片', '原高清图', '商品图片', '图片'])

                    start_row = header_row_idx + 1
                    current_row = start_row

                    for item in final_items:
                        box_display = item['box_count'] if current_carrier == "明日之星" else item['val_box_num']
                        field_writes = [
                            (col_fba, item['val_fba_no']), (col_ref, item['val_ref_id']),
                            (col_vol, item['val_vol_dim']), (col_box, box_display),
                            (col_wt, item['val_weight']), (col_hs, item['val_hs_code']),
                            (col_cn, item['val_name_cn']), (col_en, item['val_name_en']),
                            (col_qty, item['val_qty_per_box']), (col_price, item['val_unit_price']),
                            (col_mat, item['val_material']), (col_use, item['val_usage'])
                        ]

                        ws.row_dimensions[current_row].height = 65 if col_img else 28
                        for c_idx, val in field_writes:
                            if c_idx:
                                cell = ws.cell(row=current_row, column=c_idx, value=val if pd.notna(val) else "")
                                cell.alignment = center_alignment
                                cell.font = Font(name='微软雅黑', size=9)
                                cell.border = thin_border

                        # 插入商品缩略图
                        val_img_url = item['val_img_url']
                        if col_img and pd.notna(val_img_url) and str(val_img_url).startswith(('http://', 'https://')):
                            try:
                                resp = requests.get(str(val_img_url).strip(), timeout=6)
                                if resp.status_code == 200:
                                    pil_img = PILImage.open(io.BytesIO(resp.content))
                                    pil_img.thumbnail((75, 75))
                                    img_buf = io.BytesIO()
                                    pil_img.save(img_buf, format='PNG')
                                    img_buf.seek(0)
                                    xl_img = OpenpyxlImage(img_buf)
                                    xl_img.width, xl_img.height = pil_img.width, pil_img.height
                                    ws.add_image(xl_img, ws.cell(row=current_row, column=col_img).coordinate)
                                    ws.column_dimensions[openpyxl.utils.get_column_letter(col_img)].width = 16
                            except Exception:
                                pass

                        current_row += 1

                    timestamp_date = datetime.now().strftime("%Y%m%d")
                    out_filename = f"{clean_acc}_{current_carrier}_{clean_fba}-{clean_fba_wh_code}-{total_boxes}箱-{timestamp_date}.xlsx"
                    out_buf = io.BytesIO()
                    wb.save(out_buf)
                    generated_zip_files[out_filename] = out_buf.getvalue()

                    progress_bar.progress(int(40 + (current_idx / total_groups) * 55))

                # 未匹配清单
                if missing_records:
                    miss_df = pd.DataFrame(missing_records)
                    miss_buf = io.BytesIO()
                    miss_df.to_excel(miss_buf, index=False)
                    generated_zip_files[f"未匹配缺失SKU清单_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"] = miss_buf.getvalue()

                # 打包成 ZIP
                zip_buffer = io.BytesIO()
                with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                    for fn, data in generated_zip_files.items():
                        zf.writestr(fn, data)
                zip_buffer.seek(0)

                progress_bar.progress(100)
                status_text.text("✅ 所有货件发票处理完毕！")

                # 统计卡片
                st.markdown(f"""
                <div class="metric-container">
                    <div class="metric-card">
                        <div class="metric-title">识别货件数</div>
                        <div class="metric-num">{total_groups}<span class="metric-unit">个</span></div>
                    </div>
                    <div class="metric-card">
                        <div class="metric-title">成功生成发票</div>
                        <div class="metric-num">{len(generated_zip_files) - (1 if missing_records else 0)}<span class="metric-unit">份</span></div>
                    </div>
                    <div class="metric-card">
                        <div class="metric-title">跳过货件</div>
                        <div class="metric-num">{len(skipped_records)}<span class="metric-unit">个</span></div>
                    </div>
                    <div class="metric-card">
                        <div class="metric-title">未匹配 SKU</div>
                        <div class="metric-num">{len(missing_records)}<span class="metric-unit">条</span></div>
                    </div>
                </div>
                """, unsafe_allow_html=True)

                if skipped_records:
                    st.warning(f"⚠️ 有 {len(skipped_records)} 个货件因模板未配置或无法匹配而跳过。")
                if missing_records:
                    st.info(f"💡 检测到 {len(missing_records)} 条 SKU 未在采购单中匹配，已自动生成明细表放入压缩包中。")

                st.write("")

                # 居中下载按钮
                zip_filename = f"物流发票打包_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip"
                col_l, col_m, col_r = st.columns(3)
                with col_m:
                    st.download_button(
                        label="⬇️ 一键下载全部发票 (ZIP 压缩包)",
                        data=zip_buffer,
                        file_name=zip_filename,
                        mime="application/zip",
                        use_container_width=True
                    )

            except Exception as e:
                st.error(f"处理失败：{str(e)}")